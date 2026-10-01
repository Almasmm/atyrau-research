"""Build editable research reports from reviewed Markdown, without invented results.

Usage: python scripts/build_documents.py [01|02|03|04|all]
PDF conversion is separate: scripts/export_office.ps1 on Windows with Office.
"""
from pathlib import Path
import re, sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]

def clean(s,strip=True):
    s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',s)
    s=s.replace('**','').replace('`','').replace('<br>','; ')
    return s.strip() if strip else s

def runtext(p,s):
    for i,x in enumerate(re.split(r'(\*\*.*?\*\*|`[^`]+`)',s)):
        r=p.add_run(clean(x,False));r.bold=x.startswith('**')
        if x.startswith('`'):r.font.name='Consolas';r.font.size=Pt(9)

def table(doc,rows):
    t=doc.add_table(rows=0, cols=len(rows[0]));t.style='Normal Table'
    t.autofit=False
    weights=[max(8,min(28,max(len(clean(row[i])) for row in rows if len(row)>i))) for i in range(len(rows[0]))]
    total=sum(weights)
    for i,col in enumerate(t.columns):col.width=Cm(16.8*weights[i]/total)
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            if i>=len(cells):continue
            cells[i].text=clean(x);cells[i].width=Cm(16.8*weights[i]/total)
            tcPr=cells[i]._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF2' if n==0 else ('F5F7F9' if n%2==0 else 'FFFFFF'));tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge,val in [('top','65'),('bottom','65'),('left','90'),('right','90')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
                p.paragraph_format.keep_with_next=(n<len(rows)-1) if len(rows)<=8 else (n==0)
                for r in p.runs:r.font.size=Pt(8.7);r.bold=n==0;r.font.color.rgb=RGBColor(25,39,51)
        trPr=t.rows[-1]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if n==0:
            hdr=OxmlElement('w:tblHeader');trPr.append(hdr)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

def build(case):
    source=case/'Report.md'
    if not source.exists():return
    lines=source.read_text(encoding='utf-8-sig').splitlines()
    title=next((clean(x.lstrip('# ')) for x in lines if x.startswith('# ')),case.name)
    d=Document();sec=d.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.05);sec.bottom_margin=Cm(2.05);sec.left_margin=Cm(2.15);sec.right_margin=Cm(2.05)
    for s in d.styles:
        if s.type==1:
            s.font.name='Arial';s.font.color.rgb=RGBColor(20,28,35)
            pr=s.element.find(qn('w:pPr'))
            if pr is not None:
                for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for key,size in [('Title',26),('Heading 1',15),('Heading 2',12),('Heading 3',10.5)]:
        d.styles[key].font.size=Pt(size);d.styles[key].font.color.rgb=RGBColor(0,0,0)
        if key.startswith('Heading'):
            d.styles[key].paragraph_format.keep_with_next=True;d.styles[key].paragraph_format.space_before=Pt(15)
            d.styles[key].paragraph_format.space_after=Pt(6)
    d.core_properties.title=title;d.core_properties.author='';d.core_properties.subject='Четыре кейса по ИИ и машинному обучению'
    mapping={'01':'1','02':'10','03':'12','04':'15'}
    p=d.add_paragraph('ИСКУССТВЕННЫЙ ИНТЕЛЛЕКТ И МАШИННОЕ ОБУЧЕНИЕ');p.paragraph_format.space_after=Pt(48)
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(76,92,104)
    p=d.add_paragraph('Кейс '+mapping[case.name[:2]]);p.paragraph_format.space_after=Pt(14)
    for r in p.runs:r.font.size=Pt(16);r.bold=True;r.font.color.rgb=RGBColor(28,89,106)
    p=d.add_paragraph(title,style='Title');p.paragraph_format.space_after=Pt(24)
    d.add_paragraph('Преподаватель: Zhanar Oralbekova')
    d.add_paragraph('Редакция от 26 сентября 2026 года')
    p=d.add_paragraph('Данные, код, модели и результаты входят в исследовательский пакет. Происхождение данных и границы выводов указаны в отчёте.');p.paragraph_format.space_before=Pt(28)
    p=d.add_paragraph('Работа подготовлена с существенной ИИ-помощью. Подробности раскрыты в AI_USAGE_DISCLOSURE.md. Перед защитой студенту необходимо самостоятельно проверить расчёты и разобраться в методах.');p.paragraph_format.space_before=Pt(16)
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor(76,92,104)
    d.add_page_break()
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    i=0;fig=0;code=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('# '):continue
        if line.startswith('```'):
            code=not code;continue
        if code:
            p=d.add_paragraph();r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(8)
            continue
        if line.startswith('|') and i<len(lines) and re.match(r'^\|?\s*:?-',lines[i].strip()):
            rows=[[x.strip() for x in line.strip('|').split('|')]];i+=1
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]);i+=1
            table(d,rows);continue
        m=re.match(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if m:
            pth=(case/m.group(2)).resolve()
            if not pth.is_file():raise FileNotFoundError(pth)
            from PIL import Image
            with Image.open(pth) as im:w,h=im.size
            width=min(6.4,3.65*w/h)
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(pth),width=Inches(width))
            p.paragraph_format.keep_with_next=True
            fig+=1;caption=re.sub(r'^Рисунок\s+\d+[.:]?\s*','',clean(m.group(1)),flags=re.I)
            p=d.add_paragraph(f'Рисунок {fig}. '+caption,style='Caption')
            p.paragraph_format.keep_with_next=False;p.paragraph_format.space_after=Pt(10)
            for r in p.runs:r.font.size=Pt(9);r.italic=False;r.bold=True
            descr=pth.stem.replace('_',' ')
            for inline in d.inline_shapes:
                if inline._inline.docPr.get('descr') is None:inline._inline.docPr.set('descr',descr)
            continue
        m=re.match(r'^(#{2,4})\s+(.+)',line)
        if m:
            d.add_heading(clean(m.group(2)),level=len(m.group(1))-1);continue
        if line in ('---','***'):continue
        if re.match(r'^[-*]\s',line):
            p=d.add_paragraph(style='List Bullet');runtext(p,line[2:]);continue
        # The figure already has a single numbered caption. Retain its following
        # explanation without duplicating the number as a second caption.
        if re.match(r'^Рисунок\s+\d+[.]\s*',line):
            line=re.sub(r'^Рисунок\s+\d+[.]\s*','',line)
        p=d.add_paragraph();runtext(p,line)
        if re.match(r'^Таблица\s+\d+',line):
            p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(5)
            for r in p.runs:r.bold=True;r.font.size=Pt(9)
    target=case/'Report.docx';d.save(target)
    print(target.relative_to(ROOT),target.stat().st_size)

if __name__=='__main__':
    wanted=sys.argv[1] if len(sys.argv)>1 else 'all'
    for case in sorted(ROOT.glob('0[1-4]_*')):
        if wanted=='all' or case.name.startswith(wanted):build(case)
